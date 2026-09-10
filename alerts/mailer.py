"""
alerts/mailer.py — Envío de correos para alertas y reportes.
"""
import base64, io, re, smtplib, ssl
from email.mime.application import MIMEApplication
from email.mime.multipart   import MIMEMultipart
from email.mime.text        import MIMEText
from datetime                import datetime

import requests

# Por encima de esto, la tabla HTML inline se vuelve pesada (Gmail recorta
# mensajes de más de ~102KB con un "mensaje truncado") y poco legible -- se
# manda un .xlsx adjunto con el detalle completo en vez de la tabla.
MAX_INLINE_MATCHES = 300


# ── Estilos del correo ────────────────────────────────────────────────────────
_STYLE = """
<style>
  body{font-family:-apple-system,BlinkMacSystemFont,'Segoe UI',sans-serif;
       background:#0d1117;color:#e6edf3;margin:0;padding:24px}
  .wrap{max-width:680px;margin:0 auto}
  .hdr{background:#161b22;border:1px solid #30363d;border-radius:10px;
       padding:20px 24px;margin-bottom:20px}
  .hdr h1{margin:0 0 4px;font-size:18px;color:#58a6ff}
  .hdr p{margin:0;color:#8b949e;font-size:13px}
  .match{background:#161b22;border:1px solid #30363d;border-left:3px solid #3fb950;
         border-radius:8px;padding:14px 18px;margin-bottom:12px}
  .mh{display:flex;justify-content:space-between;align-items:center;margin-bottom:8px}
  .kw{background:#1f6feb33;color:#58a6ff;border:1px solid #1f6feb55;
      border-radius:4px;padding:2px 9px;font-size:12px;font-weight:700;font-family:monospace}
  .ch{color:#3fb950;font-weight:600;font-size:14px}
  .ts{color:#8b949e;font-size:12px}
  .txt{font-size:13px;color:#c9d1d9;line-height:1.6}
  .hl{color:#f0883e;font-weight:700}
  table{width:100%;border-collapse:collapse;font-size:13px}
  th{background:#21262d;color:#8b949e;padding:8px 12px;text-align:left;
     font-size:11px;text-transform:uppercase;letter-spacing:.05em}
  td{padding:8px 12px;border-bottom:1px solid #21262d;vertical-align:top}
  .ftr{text-align:center;color:#8b949e;font-size:11px;margin-top:20px}
</style>"""


def _build_smtp(cfg: dict):
    host = cfg.get('smtp_host', '')
    port = int(cfg.get('smtp_port', 587))
    user = cfg.get('smtp_user', '')
    pwd  = cfg.get('smtp_pass', '')
    tls  = cfg.get('smtp_tls', '1') == '1'
    if tls:
        ctx = ssl.create_default_context()
        srv = smtplib.SMTP(host, port, timeout=15)
        srv.ehlo(); srv.starttls(context=ctx); srv.ehlo()
    else:
        srv = smtplib.SMTP_SSL(host, port, timeout=15)
    if user:
        srv.login(user, pwd)
    return srv


def _gmail_access_token(cfg: dict) -> str:
    """Cambia el refresh_token guardado (ver gmail_oauth_setup.py) por un
    access_token de corta duración. Se pide uno nuevo en cada envío -- son
    gratis y evita tener que cachear/vencer nada."""
    r = requests.post('https://oauth2.googleapis.com/token', data={
        'client_id':     cfg.get('gmail_client_id', ''),
        'client_secret': cfg.get('gmail_client_secret', ''),
        'refresh_token': cfg.get('gmail_refresh_token', ''),
        'grant_type':    'refresh_token',
    }, timeout=15)
    r.raise_for_status()
    return r.json()['access_token']


def _send_via_gmail_api(cfg: dict, to: str, msg) -> None:
    token = _gmail_access_token(cfg)
    raw = base64.urlsafe_b64encode(msg.as_bytes()).decode('ascii')
    r = requests.post(
        'https://gmail.googleapis.com/gmail/v1/users/me/messages/send',
        headers={'Authorization': f'Bearer {token}'},
        json={'raw': raw}, timeout=20)
    r.raise_for_status()


def _send(cfg: dict, to: str, subject: str, html: str, attachment: tuple[bytes, str] = None):
    """attachment: (bytes, filename) opcional -- ej. el .xlsx del reporte
    semanal (ver send_weekly_excel_report). Sin adjunto, el mensaje sigue
    siendo un simple 'alternative' (solo HTML) como antes.

    Si hay un gmail_refresh_token guardado (ver gmail_oauth_setup.py), se
    envía por la API de Gmail sobre HTTPS en vez de SMTP -- necesario en
    redes que bloquean los puertos 25/465/587 (ver settings.html)."""
    use_api = bool(cfg.get('gmail_refresh_token'))
    sender  = cfg.get('gmail_authorized_email', '') if use_api else (cfg.get('smtp_from') or cfg.get('smtp_user', ''))
    if attachment:
        msg = MIMEMultipart('mixed')
        msg['Subject'], msg['From'], msg['To'] = subject, sender, to
        alt = MIMEMultipart('alternative')
        alt.attach(MIMEText(html, 'html', 'utf-8'))
        msg.attach(alt)
        data, filename = attachment
        part = MIMEApplication(data, Name=filename)
        part['Content-Disposition'] = f'attachment; filename="{filename}"'
        msg.attach(part)
    else:
        msg = MIMEMultipart('alternative')
        msg['Subject'], msg['From'], msg['To'] = subject, sender, to
        msg.attach(MIMEText(html, 'html', 'utf-8'))
    if use_api:
        _send_via_gmail_api(cfg, to, msg)
    else:
        srv = _build_smtp(cfg)
        srv.sendmail(sender, to, msg.as_string())
        srv.quit()


def _match_block(m: dict) -> str:
    kw   = m.get('keyword', '')
    text = m.get('matched_text', '')
    text_hl = text.replace(kw, f'<span class="hl">{kw}</span>', 1)
    ts   = str(m.get('timestamp', ''))[:19]
    ch   = m.get('channel_name', '—')
    return f"""
    <div class="match">
      <div class="mh">
        <span><span class="kw">{kw}</span>&nbsp;&nbsp;<span class="ch">{ch}</span></span>
        <span class="ts">{ts}</span>
      </div>
      <div class="txt">{text_hl}</div>
    </div>"""


def test_connection(cfg: dict, to: str) -> tuple[bool, str]:
    """Envía un correo de prueba para verificar la configuración SMTP."""
    subject = "✅ AlertaTV: prueba de configuración SMTP"
    html = f"""<!DOCTYPE html><html><head>{_STYLE}</head><body><div class="wrap">
    <div class="hdr">
      <h1>Conexión SMTP configurada correctamente</h1>
      <p>Este es un correo de prueba enviado desde AlertaTV.</p>
    </div>
    <div class="ftr">AlertaTV · {datetime.now().strftime('%d/%m/%Y %H:%M')}</div>
    </div></body></html>"""
    try:
        _send(cfg, to, subject, html)
        return True, f'Correo de prueba enviado a {to}'
    except Exception as e:
        return False, str(e)


def send_immediate(match: dict, cfg: dict) -> tuple[bool, str]:
    to = match.get('report_email', '')
    if not to:
        return False, 'Sin correo destino'
    kw  = match.get('keyword', '')
    ch  = match.get('channel_name', '—')
    sname = match.get('search_name', '')
    subject = f"\U0001f514 Alerta: «{kw}» detectado en {ch}"
    html = f"""<!DOCTYPE html><html><head>{_STYLE}</head><body><div class="wrap">
    <div class="hdr">
      <h1>Alerta detectada</h1>
      <p>Búsqueda: <strong>{sname}</strong></p>
    </div>
    {_match_block(match)}
    <div class="ftr">AlertaTV · {datetime.now().strftime('%d/%m/%Y %H:%M')}</div>
    </div></body></html>"""
    try:
        _send(cfg, to, subject, html)
        return True, f'Alerta enviada a {to}'
    except Exception as e:
        return False, str(e)


def send_threshold_alert(search: dict, count: int, window_min: int, threshold: int, cfg: dict) -> tuple[bool, str]:
    """Alerta de pico de menciones -- ver watcher.py:_check_threshold_alert.
    Distinta de send_immediate (una coincidencia puntual): aquí se reporta
    un CONTEO agregado en una ventana de tiempo, no un texto de match."""
    to = search.get('report_email', '')
    if not to:
        return False, 'Sin correo destino'
    name = search.get('name', '')
    subject = f"⚠️ Alerta de frecuencia: «{name}» — {count} en {window_min} min"
    html = f"""<!DOCTYPE html><html><head>{_STYLE}</head><body><div class="wrap">
    <div class="hdr">
      <h1>Alerta de frecuencia</h1>
      <p>Búsqueda: <strong>{name}</strong></p>
    </div>
    <div class="match">
      <div class="txt">{count} coincidencias detectadas en los últimos {window_min} minutos
      (umbral configurado: {threshold}).</div>
    </div>
    <div class="ftr">AlertaTV · {datetime.now().strftime('%d/%m/%Y %H:%M')}</div>
    </div></body></html>"""
    try:
        _send(cfg, to, subject, html)
        return True, f'Alerta de frecuencia enviada a {to}'
    except Exception as e:
        return False, str(e)


def send_weekly_excel_report(search: dict, total_matches: int, attachment_bytes: bytes,
                              attachment_name: str, cfg: dict) -> tuple[bool, str]:
    """Reporte semanal opt-in (ver watcher.py:_weekly_excel_reports) -- mismo
    Workbook que ya arma /export (alerts/excel_report.py), adjunto en vez de
    descarga directa."""
    to = search.get('report_email', '')
    if not to:
        return False, 'Sin correo destino configurado'
    name = search.get('name', '')
    subject = f"\U0001f4ca Reporte Semanal: {name} — {total_matches} coincidencias"
    html = f"""<!DOCTYPE html><html><head>{_STYLE}</head><body><div class="wrap">
    <div class="hdr">
      <h1>Reporte Semanal: {name}</h1>
      <p>Total: <strong>{total_matches}</strong> coincidencias esta semana · archivo Excel adjunto</p>
    </div>
    <div class="ftr">AlertaTV — Sistema de monitoreo TV</div>
    </div></body></html>"""
    try:
        _send(cfg, to, subject, html, attachment=(attachment_bytes, attachment_name))
        return True, f'Reporte semanal enviado a {to}'
    except Exception as e:
        return False, str(e)


def _matches_xlsx(matches: list[dict]) -> bytes:
    """.xlsx simple (sin contexto EPG/transcripción, a diferencia de
    excel_report.build_workbook) -- solo las mismas columnas de la tabla
    HTML, para reportes con demasiadas coincidencias para ir inline."""
    import openpyxl
    from openpyxl.styles import Font, PatternFill

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Coincidencias'
    ws.append(['Palabra', 'Canal', 'Fecha/Hora', 'Texto'])
    for cell in ws[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='1D4ED8')
    for m in matches:
        ws.append([
            m.get('keyword', ''),
            m.get('channel_name', ''),
            str(m.get('timestamp', ''))[:19],
            m.get('matched_text', ''),
        ])
    for col, width in zip('ABCD', (18, 24, 18, 100)):
        ws.column_dimensions[col].width = width
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def send_report(search, matches: list[dict], cfg: dict, mode: str = 'manual', summary: str = None) -> tuple[bool, str]:
    # search puede llegar como dict o como sqlite3.Row (ver watcher.py y
    # app.py:search_report) -- ambos soportan acceso por clave con [].
    to   = search['report_email']
    name = search['name']
    if not to:
        return False, 'Sin correo destino configurado'
    titles = {'daily': 'Reporte Diario', 'final': 'Reporte Final', 'manual': 'Reporte Manual'}
    title  = titles.get(mode, 'Reporte')
    subject = f"\U0001f4ca {title}: {name} — {len(matches)} coincidencias"
    attachment = None
    if len(matches) > MAX_INLINE_MATCHES:
        safe_name  = re.sub(r'[^A-Za-z0-9_-]', '_', name)
        attachment = (_matches_xlsx(matches), f'{mode}_{safe_name}.xlsx')
        table_html = f"""<div class="match">
          <div class="txt">Este reporte tiene <strong>{len(matches)}</strong> coincidencias --
          demasiadas para mostrarlas aquí. Revisa el archivo Excel adjunto para el detalle completo.</div>
        </div>"""
    else:
        rows = ''.join(f"""<tr>
          <td><span class="kw">{m.get('keyword','')}</span></td>
          <td style="color:#3fb950">{m.get('channel_name','—')}</td>
          <td style="color:#8b949e">{str(m.get('timestamp',''))[:19]}</td>
          <td style="color:#c9d1d9">{str(m.get('matched_text',''))[:140]}</td>
        </tr>""" for m in matches)
        table_html = f"""<table><tr><th>Palabra</th><th>Canal</th><th>Fecha/Hora</th><th>Texto</th></tr>
        {rows}
        </table>"""
    # Resumen ejecutivo por IA (ver rag.py:summarize_matches) -- solo en el
    # reporte diario, un bloque destacado antes de la tabla completa.
    summary_html = f"""
    <div class="hdr" style="margin-top:0;margin-bottom:20px;border-left:3px solid #58a6ff">
      <h1 style="font-size:14px;color:#c9d1d9;margin-bottom:6px">Resumen ejecutivo (generado por IA)</h1>
      <p style="font-size:13px;color:#c9d1d9;line-height:1.6;margin:0">{summary}</p>
    </div>""" if summary else ''
    html = f"""<!DOCTYPE html><html><head>{_STYLE}</head><body><div class="wrap">
    <div class="hdr">
      <h1>{title}: {name}</h1>
      <p>Total: <strong>{len(matches)}</strong> coincidencias · {datetime.now().strftime('%d/%m/%Y %H:%M')}</p>
    </div>
    {summary_html}
    {table_html}
    <div class="ftr">AlertaTV — Sistema de monitoreo TV</div>
    </div></body></html>"""
    try:
        _send(cfg, to, subject, html, attachment=attachment)
        extra = ' (Excel adjunto)' if attachment else ''
        return True, f'{title} de «{name}» enviado a {to}{extra}'
    except Exception as e:
        return False, str(e)


def send_daily_report(search, matches, cfg, summary=None):
    return send_report(search, matches, cfg, 'daily', summary=summary)

def send_final_report(search, matches, cfg):
    return send_report(search, matches, cfg, 'final')
